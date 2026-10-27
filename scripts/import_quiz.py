#!/usr/bin/env python3
"""Convert local question sources to AwayFromShorts question-bank JSON (Python 3.9+).

No app database is edited. Optional --push copies the validated file to Downloads.
JSON/CSV/TSV/marked TXT use only the standard library; XLSX needs openpyxl to read.
"""
import argparse
import csv
import hashlib
import io
import json
import math
import os
from pathlib import Path
import re
import shutil
import subprocess
import sys
import tempfile
import unicodedata

MAX_QUESTIONS = 5000
MAX_BYTES = 4 * 1024 * 1024
FIELDS = {
    "科目": "subject", "学科": "subject", "章节": "chapter", "题型": "type",
    "题目": "stem", "题干": "stem", "question": "stem", "question_text": "stem",
    "答案": "answer", "正确答案": "answer", "解析": "analysis", "解答": "analysis",
    "explanation": "analysis", "来源": "source", "题号": "id", "编号": "id",
    "选项": "options", "题库": "bank", "qtype": "type",
}
TYPES = {
    "单选": "single", "单选题": "single", "多选": "multiple", "多选题": "multiple",
    "判断": "judge", "判断题": "judge", "填空": "short", "填空题": "short",
    "解答": "solution", "解答题": "solution", "计算题": "solution", "证明题": "solution",
}
TRUE = {"正确", "对", "true", "t", "是", "√"}
FALSE = {"错误", "错", "false", "f", "否", "×"}


class ImportErrorDetail(ValueError):
    pass


def require(condition, message):
    if not condition:
        raise ImportErrorDetail(message)


def scalar(value, field):
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (int, float)):
        require(not isinstance(value, float) or math.isfinite(value), f"{field} 不能是 NaN 或无穷大")
        return str(int(value)) if isinstance(value, float) and value.is_integer() else str(value)
    require(isinstance(value, str), f"{field} 必须是文字或数字")
    value = value.strip()
    require(len(value.encode("utf-16-le")) // 2 <= 16000, f"{field} 超过 16000 字")
    return value


def subject_name(value):
    s = scalar(value, "subject")
    aliases = {"高等数学": "高数", "calculus": "高数", "math": "高数",
               "线性代数": "线代", "linear": "线代", "linear algebra": "线代"}
    s = aliases.get(s.lower(), s)
    require(0 < len(s.encode("utf-16-le")) // 2 <= 40 and not any(ord(c) < 32 or 127 <= ord(c) <= 159 for c in s),
            "科目名称须为 1–40 字，不能包含换行或控制字符")
    return s


def field_name(value):
    s = str(value).strip().lstrip("\ufeff")
    return FIELDS.get(s, s.upper() if re.fullmatch("[a-hA-H]", s) else s.lower())


def mapped(row):
    require(isinstance(row, dict), "每道题必须是对象")
    result = {}
    for key, value in row.items():
        key = field_name(key)
        require(key not in result, f"字段重复：{key}")
        result[key] = value
    return result


def choice_letters(value):
    if isinstance(value, list):
        value = "".join(scalar(v, "answer") for v in value)
    s = re.sub(r"[\s,，、;；]", "", scalar(value, "answer").upper())
    require(bool(re.fullmatch(r"[A-H]+", s)), "选择题答案请填选项字母，如 B 或 AC")
    require(len(set(s)) == len(s), "答案中有重复选项字母")
    return "".join(sorted(s))


def normalize(raw, fallback_subject, bank):
    row = mapped(raw)
    q = {"subject": subject_name(row.get("subject") or fallback_subject),
         "chapter": scalar(row.get("chapter") or "未分章", "chapter"),
         "stem": scalar(row.get("stem"), "stem")}
    require(q["stem"], "缺少题干 stem/题干")
    options = row.get("options")
    if isinstance(options, str):
        require(not options.strip() or options.strip().startswith(("[", "{")),
                "options 列请填 JSON 数组；表格也可使用 A、B、C、D 选项列")
        options = json.loads(options) if options.strip() else None
    if options is None:
        options, gap = [], False
        for key in "ABCDEFGH":
            value = scalar(row.get(key), key)
            if value:
                require(not gap, "选项须从 A 连续排列，不能跳号")
                options.append(value)
            else:
                gap = True
    elif isinstance(options, dict):
        keys = list(options)
        require(all(k in "ABCDEFGH" and len(k) == 1 for k in keys), "选项对象的键须为 A–H")
        require(set(keys) == set("ABCDEFGH"[:len(keys)]), "选项对象不能跳号")
        options = [scalar(options[k], k) for k in sorted(keys)]
    require(isinstance(options, list) and len(options) <= 8, "options 须为最多 8 个选项的数组")
    options = [scalar(v, "options") for v in options]
    require(all(options), "选项不能是空文字")
    answer = row.get("answer")
    raw_type = scalar(row.get("type"), "type")
    t = TYPES.get(raw_type, raw_type.lower())
    if not t:
        if options:
            t = "multiple" if len(choice_letters(answer)) > 1 else "single"
        else:
            t = "solution"
    require(t in {"single", "multiple", "judge", "short", "solution"}, f"未知题型：{t}")
    if t == "judge":
        options = ["正确", "错误"]
        value = scalar(answer, "answer").lower()
        answer = "A" if value in TRUE else "B" if value in FALSE else answer
    if t in {"single", "multiple", "judge"}:
        require(2 <= len(options) <= 8, "选择题需要 2–8 个选项")
        answer = choice_letters(answer)
        require(t == "multiple" or len(answer) == 1, "单选/判断题只能有一个答案")
        require(all(ord(c) - ord("A") < len(options) for c in answer), "答案字母超出选项范围")
    else:
        answer = scalar(answer, "answer")
        require(answer, "缺少答案 answer/答案；脚本不会生成或猜测答案")
    q.update(type=t, options=options, answer=answer,
             analysis=scalar(row.get("analysis"), "analysis"),
             source=scalar(row.get("source"), "source"),
             bank=scalar(row.get("bank") or bank, "bank"))
    identifier = scalar(row.get("id"), "id")
    if identifier:
        require(len(identifier.encode("utf-16-le")) // 2 <= 128 and all(
            c in "_.:-" or unicodedata.category(c)[0] in "LN" for c in identifier), "id 格式不正确（文字、数字、_.:-，最多 128 字）")
        q["id"] = identifier
    # Missing IDs stay absent: Android generates its own stable content ID.
    return q


def decode_file(path, encoding):
    require(path.stat().st_size <= 16 * 1024 * 1024, "源文件超过 16MB，请分批整理")
    raw = path.read_bytes()
    encodings = [encoding] if encoding != "auto" else ["utf-8-sig", "gb18030"]
    for enc in encodings:
        try:
            return raw.decode(enc).lstrip("\ufeff")
        except UnicodeDecodeError:
            continue
    raise ImportErrorDetail("无法解码文件，请另存为 UTF-8 或指定 --encoding")


def table_rows(rows):
    rows = iter(rows)
    header = next(rows, None)
    require(header, "表格为空")
    header = [field_name(v) if v is not None else "" for v in header]
    require(all(header) and len(set(header)) == len(header), "表头不能为空或重复")
    require("stem" in header and "answer" in header, "首行需包含 stem/题干、answer/答案 表头")
    result = []
    for number, row in enumerate(rows, 2):
        if not any(v is not None and str(v).strip() for v in row):
            continue
        require(len(row) == len(header), f"第 {number} 行列数与表头不一致")
        result.append((number, dict(zip(header, row))))
        require(len(result) <= MAX_QUESTIONS, "单个源文件最多 5000 题")
    return result


MARKER = re.compile(r"^\s*(?:\[([^\]]+)\]|【([^】]+)】|([^：:]{1,12})[：:])\s*(.*)$")


def text_rows(text):
    """Explicitly marked blocks only; never guess where a proof or formula ends."""
    rows, row, field, start = [], {}, None, 1
    for number, line in enumerate(text.splitlines(), 1):
        if re.fullmatch(r"\s*-{3,}\s*", line):
            if row:
                rows.append((start, row))
            row, field, start = {}, None, number + 1
            continue
        match = MARKER.match(line)
        key = field_name(next((g for g in match.groups()[:3] if g is not None), "")) if match else ""
        if match and key in {"subject", "chapter", "type", "stem", "answer", "analysis", "source", "id", "bank"}:
            if key == "stem" and "stem" in row:
                raise ImportErrorDetail(f"第 {number} 行：两道题之间请用 --- 分隔")
            require(key not in row, f"第 {number} 行：重复字段 {key}")
            row[key], field = match.group(4), key
            continue
        option = re.match(r"^\s*([A-H])[.、．)）]\s*(.+)$", line)
        if option and field not in {"answer", "analysis", "source"}:
            key = option.group(1)
            require(key not in row, f"第 {number} 行：选项 {key} 重复")
            row[key], field = option.group(2), key
        elif line.strip():
            require(field is not None, f"第 {number} 行：请使用 [题目]、[答案] 等标记；题间用 --- 分隔")
            row[field] += "\n" + line
        elif field:
            row[field] += "\n"
    if row:
        rows.append((start, row))
    return rows


def load_source(path, args):
    fmt = args.format if args.format != "auto" else path.suffix.lower().lstrip(".")
    fallback, name = args.subject, args.name or path.stem
    if fmt == "xlsx":
        try:
            import openpyxl
        except ImportError as exc:
            raise ImportErrorDetail("读取 XLSX 需要 openpyxl：python -m pip install openpyxl；也可另存为 CSV") from exc
        require(path.stat().st_size <= 16 * 1024 * 1024, "Excel 文件超过 16MB")
        book = openpyxl.load_workbook(path, read_only=True, data_only=False)
        try:
            sheet = book[args.sheet] if args.sheet else book.active
            rows = []
            require((sheet.max_row or 0) <= 5001 and (sheet.max_column or 0) <= 64, "Excel 最多 5000 道题、64 列，请删除多余格式区域")
            for row in sheet.iter_rows():
                require(len(rows) < 5001 and len(row) <= 64, "Excel 最多 5000 道题、64 列")
                require(not any(cell.data_type == "f" for cell in row), "Excel 含公式单元格，请复制并粘贴为值后导入")
                rows.append([cell.value for cell in row])
            return table_rows(rows), fallback, name
        finally:
            book.close()
    text = decode_file(path, args.encoding)
    if fmt == "json":
        def pairs(items):
            out = {}
            for k, v in items:
                require(k not in out, f"JSON 字段重复：{k}")
                out[k] = v
            return out
        root = json.loads(text, object_pairs_hook=pairs)
        if isinstance(root, dict):
            fallback = root.get("subject") or fallback
            name = args.name or root.get("name") or name
            root = root.get("questions")
        require(isinstance(root, list), "JSON 须为题目数组或包含 questions 数组的对象")
        return list(enumerate(root, 1)), fallback, name
    if fmt in {"csv", "tsv"}:
        return table_rows(csv.reader(io.StringIO(text, newline=""), delimiter="\t" if fmt == "tsv" else ",", strict=True)), fallback, name
    if fmt in {"txt", "md"}:
        return text_rows(text), fallback, name
    raise ImportErrorDetail("支持 JSON / CSV / TSV / XLSX / 带标记的 TXT、MD；不直接解析 PDF、Word 或图片")


def fingerprint(q):
    return json.dumps([q[k] for k in ("subject", "type", "stem", "options", "answer")], ensure_ascii=False, separators=(",", ":"))


def convert(paths, args):
    questions, seen, ids, duplicates = [], set(), {}, 0
    for path in paths:
        try:
            rows, fallback, name = load_source(path, args)
            require(0 < len(rows) <= MAX_QUESTIONS, "源文件需要 1–5000 道题")
            for number, raw in rows:
                try:
                    q = normalize(raw, fallback, name)
                    fp, identifier = fingerprint(q), q.get("id")
                    if identifier:
                        require(identifier not in ids or ids[identifier] == fp, f"id {identifier} 对应不同题目，请更换 id")
                        ids[identifier] = fp
                    if fp in seen:
                        duplicates += 1
                        continue
                    require(len(questions) < MAX_QUESTIONS, "合并后超过 5000 题，请分批导入")
                    seen.add(fp)
                    questions.append(q)
                except (ValueError, TypeError) as exc:
                    raise ImportErrorDetail(f"第 {number} 题/行：{exc}") from exc
        except (ValueError, OSError, KeyError, csv.Error) as exc:
            raise ImportErrorDetail(f"{path.name}：{exc}") from exc
    root = {"name": args.name or (paths[0].stem if len(paths) == 1 else "合并题库"), "questions": questions}
    encoded = (json.dumps(root, ensure_ascii=False, indent=2, allow_nan=False) + "\n").encode("utf-8")
    require(len(encoded) <= MAX_BYTES, "输出 JSON 超过手机版单次 4MB 上限，请分批整理")
    return root, encoded, duplicates


def write_output(path, encoded, force):
    require(force or not path.exists(), f"输出文件已存在：{path}；如需覆盖请加 --force")
    path.parent.mkdir(parents=True, exist_ok=True)
    descriptor, temporary = tempfile.mkstemp(prefix=".afs-quiz-", dir=str(path.parent))
    try:
        with os.fdopen(descriptor, "wb") as stream:
            stream.write(encoded)
            stream.flush()
            os.fsync(stream.fileno())
        if force:
            os.replace(temporary, path)
        else:
            # Exclusive creation prevents another process's output being overwritten.
            os.link(temporary, path)
    finally:
        if os.path.exists(temporary):
            os.unlink(temporary)


def push_to_phone(path, args):
    adb = args.adb or os.environ.get("AFS_ADB") or shutil.which("adb")
    if not adb:
        candidate = Path.home() / ".afs-android/sdk/platform-tools/adb.exe"
        adb = str(candidate) if candidate.is_file() else None
    require(adb, "找不到 adb；请指定 --adb 路径。JSON 已生成，可手动传到手机")
    prefix = [adb] + (["-s", args.serial] if args.serial else [])
    status = subprocess.run(prefix + ["get-state"], capture_output=True, text=True, timeout=20)
    require(status.returncode == 0 and status.stdout.strip() == "device", "手机未连接/未授权，或连接了多台设备；多设备请用 --serial")
    folder = "/sdcard/Download/AFS-Quiz"
    subprocess.run(prefix + ["shell", "mkdir", "-p", folder], check=True, timeout=20)
    # A fixed ASCII destination avoids shell/path interpretation of user-controlled names.
    destination = folder + "/bank-" + hashlib.sha256(path.read_bytes()).hexdigest()[:12] + ".json"
    subprocess.run(prefix + ["push", str(path), destination], check=True, timeout=120)
    print(f"已传至手机：{destination}\n请在刷题页点“导入题库”，从下载目录选择此文件并确认预览。")


def template(fmt):
    q = {"subject": "高数", "chapter": "极限与连续", "type": "single",
         "stem": r"求 $\lim_{x\to0}\frac{\sin x}{x}$。", "options": ["$0$", "$1$", "$-1$", "不存在"],
         "answer": "B", "analysis": "由基本重要极限，答案为 $1$。", "source": "自编示例"}
    if fmt == "json":
        return (json.dumps({"name": "题库模板", "questions": [q]}, ensure_ascii=False, indent=2) + "\n").encode("utf-8")
    if fmt == "txt":
        return ("[科目] 高数\n[章节] 极限与连续\n[题型] 单选\n[题目] " + q["stem"] +
                "\nA. $0$\nB. $1$\nC. $-1$\nD. 不存在\n[答案] B\n[解析] " + q["analysis"] +
                "\n[来源] 自编示例\n---\n[科目] 线代\n[章节] 矩阵\n[题型] 解答题\n[题目] 求单位矩阵 $I$ 的逆。\n[答案] $I$\n[解析] $II=I$。\n").encode("utf-8")
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(["科目", "章节", "题型", "题干", "A", "B", "C", "D", "答案", "解析", "来源"])
    writer.writerow([q["subject"], q["chapter"], q["type"], q["stem"], *q["options"], q["answer"], q["analysis"], q["source"]])
    return stream.getvalue().encode("utf-8-sig")


def main(argv=None):
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, "reconfigure"):
            stream.reconfigure(encoding="utf-8")
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("inputs", nargs="*", type=Path, help="一个或多个源文件，按顺序合并去重")
    parser.add_argument("-o", "--output", type=Path, help="输出 UTF-8 JSON 路径")
    parser.add_argument("--subject", default="高数", help="未指定科目时的默认值（支持自定义科目）")
    parser.add_argument("--name", help="输出题库名")
    parser.add_argument("--format", choices=["auto", "json", "csv", "tsv", "xlsx", "txt", "md"], default="auto")
    parser.add_argument("--encoding", default="auto", help="文本源编码；auto 尝试 UTF-8、GB18030")
    parser.add_argument("--sheet", help="Excel 工作表名称，默认活动表")
    parser.add_argument("--template", choices=["json", "csv", "txt"], help="创建可填写的模板")
    parser.add_argument("--check", action="store_true", help="只检查、统计，不写文件")
    parser.add_argument("--force", action="store_true", help="允许覆盖输出文件（仍禁止覆盖输入源文件）")
    parser.add_argument("--push", action="store_true", help="生成后用 adb 复制到手机下载目录")
    parser.add_argument("--adb", help="adb 可执行文件路径")
    parser.add_argument("--serial", help="指定 adb 设备序列号")
    args = parser.parse_args(argv)
    try:
        require(not args.push or not args.check, "--check 不能与 --push 同用")
        require(not args.template or not args.inputs, "创建模板时不应指定输入文件")
        require(not args.template or not args.check, "模板创建不能使用 --check")
        require(args.output is not None or args.check, "请用 -o 指定输出文件，或用 --check 只检查")
        if args.template:
            encoded = template(args.template)
        else:
            require(args.inputs, "请指定题库源文件，或用 --template 创建模板")
            root, encoded, duplicates = convert(args.inputs, args)
            counts = {}
            for q in root["questions"]:
                counts[q["subject"]] = counts.get(q["subject"], 0) + 1
            print("校验通过：" + "，".join(f"{k} {v} 题" for k, v in counts.items()))
            print(f"合计 {len(root['questions'])} 题，跳过重复 {duplicates} 题，JSON {len(encoded):,} 字节")
        if args.check:
            return 0
        require(all(args.output.resolve() != p.resolve() for p in args.inputs), "输出不能覆盖源文件，请换一个 -o 路径")
        write_output(args.output, encoded, args.force)
        print(f"已生成：{args.output.resolve()}")
        if args.push:
            push_to_phone(args.output, args)
        return 0
    except (ValueError, OSError, LookupError, subprocess.SubprocessError) as exc:
        print(f"错误：{exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
